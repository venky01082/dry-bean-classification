import os
import json
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
blank_layout = prs.slide_layouts[6]

DARK_BG = RGBColor(15, 23, 42)
CARD_BG = RGBColor(30, 41, 59)
PRIMARY = RGBColor(16, 185, 129)
ACCENT = RGBColor(2, 132, 199)
TEXT_WHITE = RGBColor(255, 255, 255)
TEXT_MUTED = RGBColor(148, 163, 184)

def add_header(slide, title_text, category="B.TECH AI & DATA SCIENCE"):
    # Header box
    tb = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.7), Inches(1.1))
    tf = tb.text_frame
    tf.word_wrap = True
    
    p0 = tf.paragraphs[0]
    p0.text = category.upper()
    p0.font.size = Pt(11)
    p0.font.bold = True
    p0.font.color.rgb = PRIMARY
    
    p1 = tf.add_paragraph()
    p1.text = title_text
    p1.font.size = Pt(24)
    p1.font.bold = True
    p1.font.color.rgb = DARK_BG

# --- Slide 1: Title ---
s1 = prs.slides.add_slide(blank_layout)
tb1 = s1.shapes.add_textbox(Inches(1.0), Inches(1.5), Inches(11.3), Inches(4.5))
tf1 = tb1.text_frame
tf1.word_wrap = True

p = tf1.paragraphs[0]
p.text = "🌱 BeanAI — Intelligent Dry Bean Classification"
p.font.size = Pt(32)
p.font.bold = True
p.font.color.rgb = DARK_BG

p = tf1.add_paragraph()
p.text = "Multiclass Crop Cultivar Prediction Using Machine Learning & Interactive Web Dashboard"
p.font.size = Pt(18)
p.font.color.rgb = ACCENT
p.space_after = Pt(24)

p = tf1.add_paragraph()
p.text = "Course: B.Tech in Artificial Intelligence & Data Science"
p.font.size = Pt(14)
p.font.bold = True
p.font.color.rgb = DARK_BG

p = tf1.add_paragraph()
p.text = "Team Members & IDs:"
p.font.size = Pt(13)
p.font.bold = True
p.font.color.rgb = DARK_BG

members = [
    ("24EU01007", "Data Engineering, EDA & Data Preprocessing"),
    ("24EU01046", "ML Modeling, Hyperparameter Tuning & Evaluation"),
    ("24EU01055", "FastAPI Backend, React Dashboard & Integration")
]
for sid, role in members:
    p = tf1.add_paragraph()
    p.text = f"• {sid} — {role}"
    p.font.size = Pt(12)
    p.font.color.rgb = RGBColor(51, 65, 85)

# --- Slide 2: Problem Statement ---
s2 = prs.slides.add_slide(blank_layout)
add_header(s2, "Problem Statement & Agricultural Motivation")
tb2 = s2.shapes.add_textbox(Inches(0.8), Inches(1.8), Inches(11.7), Inches(5.0))
tf2 = tb2.text_frame
tf2.word_wrap = True
points2 = [
    "Industrial Context: Turkey is among the world's leading producers and exporters of dry edible beans (Phaseolus vulgaris L.).",
    "Market Challenge: 7 distinct commercial varieties command significantly different market values, cooking properties, and culinary applications.",
    "The Problem: Varieties often look identical to the naked eye. Manual seed grading is subjective, labor-intensive, error-prone, and inconsistent.",
    "System Objective: Develop an automated, objective Machine Learning system to classify bean varieties from 16 physical geometric measurements.",
    "Outcome: High-accuracy multiclass classifier (>92%) deployed via a real-time FastAPI service and interactive React frontend."
]
for pt in points2:
    p = tf2.add_paragraph()
    p.text = "• " + pt
    p.font.size = Pt(15)
    p.space_after = Pt(14)
    p.font.color.rgb = RGBColor(30, 41, 59)

# --- Slide 3: Dataset & Data Source ---
s3 = prs.slides.add_slide(blank_layout)
add_header(s3, "Dataset Source & Class Characteristics")
tb3 = s3.shapes.add_textbox(Inches(0.8), Inches(1.8), Inches(11.7), Inches(5.0))
tf3 = tb3.text_frame
tf3.word_wrap = True
points3 = [
    "Dataset Origin: UCI Machine Learning Repository (Koklu & Ozkan, 2020).",
    "Sample Dimensions: 13,611 high-resolution images segmented into 16 numerical geometric features.",
    "Zero Missing Values: Complete, verified dataset with 0 missing entries.",
    "Target: 7 Cultivar Classes:",
    "   1. DERMASON: 3,546 samples (26.05%) — Most popular commercial flat-oval white bean.",
    "   2. SIRA: 2,636 samples (19.37%) — Small-medium compact white bean.",
    "   3. SEKER: 2,027 samples (14.89%) — Spherical bean with highest roundness (0.958).",
    "   4. HOROZ: 1,928 samples (14.16%) — Rooster bean, elongated cylindrical profile.",
    "   5. CALI: 1,630 samples (11.98%) — Large kidney-shaped variety.",
    "   6. BARBUNYA: 1,322 samples (9.71%) — Mottled borlotti pinto-type bean.",
    "   7. BOMBAY: 522 samples (3.84%) — Massive white runner bean (>150,000 px² area)."
]
for pt in points3:
    p = tf3.add_paragraph()
    p.text = pt
    p.font.size = Pt(13)
    p.space_after = Pt(6)
    p.font.color.rgb = RGBColor(30, 41, 59)

# --- Slide 4: EDA & Visualizations ---
s4 = prs.slides.add_slide(blank_layout)
add_header(s4, "Exploratory Data Analysis & Feature Distributions")
if os.path.exists("reports/class_distribution.png"):
    s4.shapes.add_picture("reports/class_distribution.png", Inches(0.8), Inches(1.8), width=Inches(5.6))
if os.path.exists("reports/feature_boxplots.png"):
    s4.shapes.add_picture("reports/feature_boxplots.png", Inches(6.8), Inches(1.8), width=Inches(5.7))

# --- Slide 5: Data Preprocessing & Leakage Prevention ---
s5 = prs.slides.add_slide(blank_layout)
add_header(s5, "Data Preprocessing & Leakage Prevention")
tb5 = s5.shapes.add_textbox(Inches(0.8), Inches(1.8), Inches(11.7), Inches(5.0))
tf5 = tb5.text_frame
tf5.word_wrap = True
points5 = [
    "Stratified 80/20 Train/Test Split: 10,888 training samples, 2,723 test samples.",
    "Strict Leakage Prevention: All feature scaling (StandardScaler) was encapsulated within an sklearn Pipeline. Scaler parameters (mean, std) were fitted strictly on training folds.",
    "Feature Integrity: Validated exact 16 predictor columns in precise order required by the mathematical model.",
    "Collinearity Analysis: Correlation matrix revealed r > 0.99 between Area and ConvexArea; dimensionless Shape Factors (1-4) were retained for non-scale morphology separation."
]
for pt in points5:
    p = tf5.add_paragraph()
    p.text = "• " + pt
    p.font.size = Pt(15)
    p.space_after = Pt(14)
    p.font.color.rgb = RGBColor(30, 41, 59)

# --- Slide 6: Three ML Models Compared ---
s6 = prs.slides.add_slide(blank_layout)
add_header(s6, "Three ML Architectures Evaluated")
tb6 = s6.shapes.add_textbox(Inches(0.8), Inches(1.8), Inches(11.7), Inches(5.0))
tf6 = tb6.text_frame
tf6.word_wrap = True
points6 = [
    "1. Logistic Regression (Linear Reference Baseline):",
    "   - Softmax multiclass classification with L2 regularization and StandardScaler.",
    "   - Fast, interpretable benchmark across 7 classes.",
    "2. Random Forest Classifier (Non-linear Ensemble):",
    "   - 100-200 bagging decision trees with randomized feature splits.",
    "   - Scale-invariant, captures non-linear feature interactions, calculates Gini feature importance.",
    "3. Support Vector Machine (SVM - Margin Based):",
    "   - Radial Basis Function (RBF) kernel mapping continuous measurements into infinite-dimensional space.",
    "   - Standardized features with calibrated probability estimates enabled."
]
for pt in points6:
    p = tf6.add_paragraph()
    p.text = pt
    p.font.size = Pt(14)
    p.space_after = Pt(8)
    p.font.color.rgb = RGBColor(30, 41, 59)

# --- Slide 7: Hyperparameter Tuning via Stratified 5-Fold CV ---
s7 = prs.slides.add_slide(blank_layout)
add_header(s7, "Hyperparameter Tuning (Stratified 5-Fold Cross Validation)")
tb7 = s7.shapes.add_textbox(Inches(0.8), Inches(1.8), Inches(11.7), Inches(5.0))
tf7 = tb7.text_frame
tf7.word_wrap = True
points7 = [
    "Search Strategy: Exhaustive GridSearchCV with 5-fold Stratified K-Fold cross-validation on 10,888 training rows.",
    "Optimization Metric: Macro F1-Score (guarantees equal weighting across minority classes like Bombay).",
    "Tuning Results:",
    "   • Logistic Regression: C in [0.1, 1.0, 10.0], solver in ['lbfgs', 'saga'] -> Best: C=10.0, solver='saga' (CV F1: 0.9368)",
    "   • Random Forest: n_estimators in [100, 200], max_depth in [15, 25, None] -> Best: max_depth=15, n=100 (CV F1: 0.9358)",
    "   • Support Vector Machine: C in [1.0, 10.0], gamma in ['scale', 'auto'] -> Best: C=10.0, kernel='rbf', gamma='scale' (CV F1: 0.9443)"
]
for pt in points7:
    p = tf7.add_paragraph()
    p.text = pt
    p.font.size = Pt(14)
    p.space_after = Pt(10)
    p.font.color.rgb = RGBColor(30, 41, 59)

# --- Slide 8: Model Comparison & Final Evaluation ---
s8 = prs.slides.add_slide(blank_layout)
add_header(s8, "Model Comparison & Final Test Set Evaluation")
tb8 = s8.shapes.add_textbox(Inches(0.8), Inches(1.8), Inches(11.7), Inches(5.0))
tf8 = tb8.text_frame
tf8.word_wrap = True
points8 = [
    "Evaluated on Held-out Test Set (2,723 samples, 20%):",
    "",
    "   Model                             Accuracy      Macro Precision    Macro Recall    Macro F1    Train Time",
    "   ------------------------------------------------------------------------------------------------------",
    "   Support Vector Machine (SVM) 🏆   92.43%        93.85%             93.42%          0.9361      20.8s",
    "   Logistic Regression               92.14%        93.60%             93.10%          0.9332      47.1s",
    "   Random Forest                     91.88%        93.48%             92.83%          0.9313      56.4s",
    "",
    "Key Findings:",
    "• Support Vector Machine achieved highest test accuracy (92.43%) and Macro F1 (0.9361).",
    "• Average inference time per bean measurement: 0.16 milliseconds."
]
for pt in points8:
    p = tf8.add_paragraph()
    p.text = pt
    p.font.size = Pt(13)
    p.font.color.rgb = RGBColor(30, 41, 59)

# --- Slide 9: Model Persistence & System Architecture ---
s9 = prs.slides.add_slide(blank_layout)
add_header(s9, "Model Persistence & Production Architecture")
tb9 = s9.shapes.add_textbox(Inches(0.8), Inches(1.8), Inches(11.7), Inches(5.0))
tf9 = tb9.text_frame
tf9.word_wrap = True
points9 = [
    "Pipeline Serialization: Saved the entire fitted pipeline (StandardScaler + SVC) using Joblib into models/dry_bean_pipeline.joblib.",
    "Reload Verification: Loaded pipeline into backend memory on startup; verified predictions match original training output exactly.",
    "FastAPI Backend: REST API with Pydantic type validation enforcing exact 16-feature schema.",
    "   • GET /health — Health check & model status",
    "   • GET /model-info — Model parameters and metrics",
    "   • POST /predict — Real-time inference with probability distribution across all 7 classes",
    "   • GET /samples — Curated real bean presets"
]
for pt in points9:
    p = tf9.add_paragraph()
    p.text = pt
    p.font.size = Pt(14)
    p.space_after = Pt(10)
    p.font.color.rgb = RGBColor(30, 41, 59)

# --- Slide 10: Interactive Frontend Dashboard ---
s10 = prs.slides.add_slide(blank_layout)
add_header(s10, "Interactive Web Application (BeanAI)")
tb10 = s10.shapes.add_textbox(Inches(0.8), Inches(1.8), Inches(11.7), Inches(5.0))
tf10 = tb10.text_frame
tf10.word_wrap = True
points10 = [
    "Frontend Stack: React 19 + Vite + Recharts + Lucide Icons.",
    "Core Features:",
    "   1. Dashboard: Dataset metrics (13,611 rows, 16 features), class distributions bar/pie charts, and cultivar catalogue.",
    "   2. Bean Prediction Form: 16 organized input fields with 1-click real sample presets for all 7 bean varieties.",
    "   3. Interactive Probability Chart: Recharts horizontal bar chart showing class probability confidence.",
    "   4. Model Performance Page: Side-by-side comparison bar chart, interactive confusion matrix, and feature importances.",
    "   5. About & Team Page: Complete documentation, team roles, and scientific citations."
]
for pt in points10:
    p = tf10.add_paragraph()
    p.text = pt
    p.font.size = Pt(14)
    p.space_after = Pt(10)
    p.font.color.rgb = RGBColor(30, 41, 59)

# --- Slide 11: Real-World Applications & Limitations ---
s11 = prs.slides.add_slide(blank_layout)
add_header(s11, "Real-World Applications & Technical Limitations")
tb11 = s11.shapes.add_textbox(Inches(0.8), Inches(1.8), Inches(11.7), Inches(5.0))
tf11 = tb11.text_frame
tf11.word_wrap = True
points11 = [
    "Industrial Applications:",
    "   • High-speed optical sorters in agricultural grain packaging plants.",
    "   • Quality control inspection and cultivar certification for seed export agencies.",
    "   • Pricing and purity verification for automated grain elevators.",
    "Limitations & Scope:",
    "   • Feature dependency on optical segmentation quality (camera lighting and calibration).",
    "   • Misclassification bottleneck between Sira and Dermason (~7% overlap) due to biological morphology proximity.",
    "   • Future Work: Integrating convolutional neural networks (CNNs) for direct image inference alongside tabular physical traits."
]
for pt in points11:
    p = tf11.add_paragraph()
    p.text = pt
    p.font.size = Pt(14)
    p.space_after = Pt(10)
    p.font.color.rgb = RGBColor(30, 41, 59)

# --- Slide 12: Conclusion & Q&A ---
s12 = prs.slides.add_slide(blank_layout)
add_header(s12, "Conclusion & Project Defense")
tb12 = s12.shapes.add_textbox(Inches(0.8), Inches(1.8), Inches(11.7), Inches(5.0))
tf12 = tb12.text_frame
tf12.word_wrap = True
points12 = [
    "Summary of Achievements:",
    "   ✓ Complete ML lifecycle: Acquisition -> EDA -> Preprocessing -> Modeling -> Tuning -> Persistence -> Deployment.",
    "   ✓ Evaluated 3 candidate models; optimized champion SVM achieving 92.43% test accuracy and 0.9361 Macro F1.",
    "   ✓ Production-grade FastAPI backend with zero-retraining inference.",
    "   ✓ Modern, responsive React dashboard with 1-click real bean presets and probability visualizations.",
    "   ✓ Comprehensive automated test suite (9/9 unit tests passing).",
    "",
    "Thank you! We welcome questions from the faculty panel."
]
for pt in points12:
    p = tf12.add_paragraph()
    p.text = pt
    p.font.size = Pt(14)
    p.space_after = Pt(8)
    p.font.color.rgb = RGBColor(30, 41, 59)

out_pptx = "reports/Dry_Bean_Classification_Presentation.pptx"
prs.save(out_pptx)
print(f"PowerPoint presentation generated at: {out_pptx}")
